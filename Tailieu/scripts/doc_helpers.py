import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_FAMILY = "Times New Roman"
COLOR_PRIMARY = RGBColor(31, 73, 125)     # Deep Navy Blue #1F497D
COLOR_SECONDARY = RGBColor(192, 0, 0)     # Academic Dark Red #C00000
COLOR_TEXT = RGBColor(38, 38, 38)         # Dark Gray Charcoal
COLOR_MUTED = RGBColor(100, 100, 100)

def set_page_setup(doc):
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.787)      # 2.0 cm
        section.bottom_margin = Inches(0.787)   # 2.0 cm
        section.left_margin = Inches(1.181)     # 3.0 cm
        section.right_margin = Inches(0.787)    # 2.0 cm

def format_run(run, font_name=FONT_FAMILY, size_pt=12, bold=False, italic=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

def add_p(doc, text="", bold_prefix="", italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        format_run(r_pre, bold=True)
    if text:
        r_txt = p.add_run(text)
        format_run(r_txt, italic=italic)
    return p

def add_bullet(doc, text="", bold_prefix="", space_after=3):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        format_run(r_pre, bold=True)
    if text:
        r_txt = p.add_run(text)
        format_run(r_txt)
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    format_run(r, size_pt=14, bold=True, color_rgb=COLOR_PRIMARY)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    format_run(r, size_pt=13, bold=True, color_rgb=COLOR_PRIMARY)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    format_run(r, size_pt=12, bold=True, italic=True, color_rgb=COLOR_TEXT)
    return p

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def create_styled_table(doc, headers, data, col_widths=None, header_bg="1F497D", stripe=True):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Apply table borders
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="B0C4DE"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1F497D"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E6ED"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

    # Header Row
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    header_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, header_bg)
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(h_text)
        format_run(r, size_pt=10.5, bold=True, color_rgb=RGBColor(255, 255, 255))

    # Data Rows
    for row_idx, row_values in enumerate(data):
        row_num = row_idx + 1
        tr = table.rows[row_num]._tr.get_or_add_trPr()
        tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_color = "F9FBFD" if (stripe and row_idx % 2 == 1) else "FFFFFF"

        for col_idx, val in enumerate(row_values):
            cell = table.cell(row_num, col_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if bg_color != "FFFFFF":
                set_cell_shading(cell, bg_color)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            
            # Align center for short codes or IDs
            val_str = str(val).strip()
            if col_idx == 0 and len(val_str) <= 6:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif val_str in ['Pass', 'Fail', 'PK', 'FK', 'NN', 'UQ', 'ACTIVE', 'PAID']:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
            r = p.add_run(val_str)
            format_run(r, size_pt=10)

    # Column Widths
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)

    # Add spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)
    return table

def add_figure(doc, img_path, caption="", width=Inches(6.0)):
    if not os.path.exists(img_path):
        print(f"Warning: Image not found at {img_path}")
        return None
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=width)

    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(caption)
        format_run(r_cap, size_pt=10.5, italic=True, color_rgb=COLOR_MUTED)
        return p_cap
    return p_img
