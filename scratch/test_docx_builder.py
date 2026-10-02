import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Pt, Inches
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

test_doc_path = r"d:\Detai12_QLCH\scratch\test_output.docx"
doc = docx.Document(r"d:\Detai12_QLCH\Tailieu\05_GenAI_SoftwareDevelopment_functional-testing.docx.bak")

def set_cell_properties(cell, align='left', font_size_pt=13, bold=False, text=''):
    p = cell.paragraphs[0]
    p.text = ''
    pPr = p._p.get_or_add_pPr()
    spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="120" w:line="240" w:lineRule="auto"/>')
    pPr.append(spacing)
    if align in ('center', 'right'):
        jc = parse_xml(f'<w:jc {nsdecls("w")} w:val="{align}"/>')
        pPr.append(jc)
        
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="ffffff"/>')
    tcPr.append(shd)
    vAlign = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>')
    tcPr.append(vAlign)
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold

tbl = doc.tables[2]
# Test adding a row
new_row = tbl.add_row()
for c_idx, cell in enumerate(new_row.cells):
    set_cell_properties(cell, align='center' if c_idx == 0 else 'left', text=f'Test {c_idx}')

doc.save(test_doc_path)

# Reopen and check
doc_check = docx.Document(test_doc_path)
print("Successfully saved and reloaded! Total rows in Table 2:", len(doc_check.tables[2].rows))
