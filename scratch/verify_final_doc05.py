import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc_path = r"d:\Detai12_QLCH\Tailieu\05_GenAI_SoftwareDevelopment_functional-testing.docx"
doc = docx.Document(doc_path)

print("=== VERIFYING FINALIZED DOCX ===")
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"  P{i:02d}: '{p.text[:70]}' (runs: {len(p.runs)})")

print("\n--- TABLES INSPECTION ---")
for idx, tbl in enumerate(doc.tables):
    print(f"Table {idx}: {len(tbl.rows)} rows x {len(tbl.columns)} cols")
    # inspect header
    h_texts = [c.text.strip().replace('\n', ' ') for c in tbl.rows[0].cells]
    print(f"  Header: {h_texts}")
    # inspect first data row
    if len(tbl.rows) > 1:
        d1_texts = [c.text.strip().replace('\n', ' ') for c in tbl.rows[1].cells]
        print(f"  Row 1 : {d1_texts}")
        # inspect first cell run font
        r = tbl.rows[1].cells[0].paragraphs[0].runs[0]
        print(f"  Row 1 Font: {r.font.name}, {r.font.size.pt if r.font.size else 'None'}pt, bold={r.font.bold}")
    # inspect last row
    if len(tbl.rows) > 2:
        last_texts = [c.text.strip().replace('\n', ' ') for c in tbl.rows[-1].cells]
        print(f"  Last  : {last_texts}")

print("\n=== VERIFICATION SUCCESSFUL ===")
